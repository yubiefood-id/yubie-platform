"""Offline workbook preparation/conversion. Never connects to a database or PSP."""
import argparse
import datetime
import hashlib
import json
import pathlib
import re
import subprocess
import sys
import zipfile
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

CATALOG = [
    ('flour','250g','YBF-250',15,None,1), ('flour','1kg','YBF-1000',5,None,1),
    ('shake','sachet','YBS-045',20,None,1), ('shake','pack-7','YBS-7X45',None,'YBS-045',7),
    ('ppang','mini','YBP-MINI-035',15,None,1), ('ppang','maxi','YBP-MAXI',5,None,1),
    ('ppang','mini-pack-5','YBP-MINI-5',None,'YBP-MINI-035',5),
    ('mie','single','YBM-075',20,None,1), ('mie','pack-3','YBM-3X75',None,'YBM-075',3),
]
SKU = {row[2]:row for row in CATALOG}
SHEETS = {
 'Pengesahan Intake':['Kolom','Nilai','Petunjuk'],
 'Intake SKU':['SKU','Status verifikasi','Mode fisik: independent/derived','SKU komponen','Jumlah komponen','Usulan (bukan stok)','Jumlah fisik terkonfirmasi','Versi kebijakan','Referensi persetujuan kemasan'],
 'Intake Lot QC':['SKU','Kode lot asli','Jumlah tersedia fisik','Saldo tersedia DB sebelum hitung','UOM','Tanggal kedaluwarsa YYYY-MM-DD','Waktu hitung UTC ISO','Status lot','Referensi rilis QC','Nama penanggung jawab QC','Status verifikasi'],
 'Intake Pengiriman':['Status verifikasi','Kota/kabupaten (satu per baris)','Kebijakan','Ongkir IDR','Layanan dan cakupan','Penanggung biaya','Mulai UTC ISO','Berakhir UTC ISO','Referensi persetujuan','Pemberi persetujuan','Referensi bukti','Waktu verifikasi UTC ISO','Cakupan disetujui: whole_city'],
 'Staf dan Finance':['Keputusan yang belum tersedia','Nilai asli','Petunjuk'],
}

def check_archive(path):
    path=pathlib.Path(path)
    if path.suffix.lower()!='.xlsx':raise ValueError('Hanya .xlsx; proposal JSON NOT_FOR_IMPORT bukan intake')
    if path.stat().st_size>10_000_000:raise ValueError('Workbook terlalu besar')
    with zipfile.ZipFile(path) as archive:
        if sum(i.file_size for i in archive.infolist())>100_000_000:raise ValueError('Archive terlalu besar')
        if any('vbaProject' in i.filename or 'externalLinks/' in i.filename for i in archive.infolist()):raise ValueError('Macro/external links tidak diterima')

def prepare(source, destination):
    source,destination=pathlib.Path(source),pathlib.Path(destination)
    if source.resolve()==destination.resolve() or destination.exists():raise ValueError('Gunakan berkas keluaran baru; sumber tidak boleh ditimpa')
    check_archive(source); workbook=load_workbook(source,keep_links=False)
    if any(name in workbook for name in SHEETS):raise ValueError('Worksheet intake sudah ada; tidak ditimpa')
    for name,header in SHEETS.items():workbook.create_sheet(name).append(header)
    meta=workbook['Pengesahan Intake']
    for row in [
        ('status','PENDING','Isi VERIFIED hanya setelah bukti fisik sah tersedia'),
        ('evidenceRef','PENDING','Referensi dokumen hitung fisik dan QC asli'),
        ('verifiedBy','PENDING','Nama penanggung jawab hitung'),
        ('verifiedAt','PENDING','Waktu verifikasi asli dalam UTC ISO'),
        ('countedAt','PENDING','Waktu hitung asli UTC ISO; berlaku 15 menit saat apply'),
        ('purpose','initial_inventory','20 per keluarga; hitung ulang harian gunakan operational_count'),
        ('actor','PENDING','Akun staf; server memakai identitas login'),
        ('reason','PENDING','Alasan dan referensi hitung asli'),
    ]:meta.append(row)
    for family,size,sku,proposed,component,quantity in CATALOG:
        workbook['Intake SKU'].append([sku,'PENDING','derived' if component else 'independent',component or sku,quantity,proposed,None,1,'PENDING'])
    for field,hint in [
        ('Google client ID khusus staf','Client terpisah dari login pelanggan; bukan API key'),
        ('Daftar subject Google dan peran staf','Named accounts; pemeriksa QC/Finance harus berbeda dari pembuat'),
        ('Referensi persetujuan akses staf','Bukti persetujuan pemilik yang asli'),
        ('Penanggung jawab refund dan eskalasi Finance','Nama penanggung jawab dan referensi prosedur'),
        ('Bukti acceptance pembayaran positif','Referensi sesi/callback/rekonsiliasi tersanitasi; jangan tempel token atau signature'),
        ('Persetujuan pilot uang nyata','Maksimum nominal/jumlah transaksi, periode UTC, pembeli milik merchant, pelaksana dan referensi approval'),
    ]:workbook['Staf dan Finance'].append([field,'PENDING',hint])
    # Deliberately no physical lots or shipping approval rows are invented.
    for name in SHEETS:
        sheet=workbook[name];sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='633490')
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width=min(44,max(20,len(str(column[0].value or ''))+2))
            for cell in column:cell.alignment=Alignment(vertical='top',wrap_text=True)
    workbook.save(destination)

def text(value,field):
    if not isinstance(value,str) or not value.strip():raise ValueError(field+': belum diisi')
    value=value.strip()
    if value.startswith('='):raise ValueError(field+': formula tidak diterima')
    if re.search(r'(^|[\s_-])(PENDING|PROPOSED|BELUM|USULAN|NOT_FOR_IMPORT)($|[\s_-])',value,re.I):raise ValueError(field+': PENDING/belum VERIFIED')
    return value

def number(value,field,minimum=0):
    if isinstance(value,bool) or not isinstance(value,(int,float)) or int(value)!=value or value<minimum or value>2_147_483_647:raise ValueError(field+': harus bilangan bulat sah')
    return int(value)

def timestamp(value,field):
    value=text(value,field)
    try: parsed=datetime.datetime.fromisoformat(value.replace('Z','+00:00'))
    except ValueError:raise ValueError(field+': UTC ISO diperlukan') from None
    if parsed.tzinfo is None or parsed.utcoffset()!=datetime.timedelta(0):raise ValueError(field+': timezone UTC diperlukan')
    return parsed.isoformat().replace('+00:00','Z')

def rows(workbook,name):
    if name not in workbook:raise ValueError(name+': worksheet intake belum ada; jalankan prepare')
    sheet=workbook[name]
    if sheet.max_row>2000:raise ValueError('Terlalu banyak baris')
    result=[]
    for index,row in enumerate(sheet.iter_rows(min_row=2,values_only=True),2):
        if not any(v is not None for v in row):continue
        if any(isinstance(v,str) and v.startswith('=') for v in row):raise ValueError(f'{name} baris {index}: formula tidak diterima')
        result.append(row)
    return result

def convert(path,scope):
    path=pathlib.Path(path);check_archive(path)
    workbook=load_workbook(path,data_only=False,keep_links=False)
    digest=hashlib.sha256(path.read_bytes()).hexdigest()
    if scope=='shipping':
        shipping=rows(workbook,'Intake Pengiriman')
        if not shipping:raise ValueError('Pengiriman: cakupan dan persetujuan belum diberikan')
        policies=[]
        for row in shipping:
            status,city,policy,fee,coverage,cost,start,end,reference,approver,evidence,verified=row[:12]
            if text(status,'Pengiriman status')!='VERIFIED':raise ValueError('Pengiriman harus VERIFIED')
            if policy!='free_promotional' or number(fee,'Ongkir')!=0:raise ValueError('Schema saat ini hanya free_promotional IDR 0 dengan persetujuan sah')
            if len(row)<13 or text(row[12],'Cakupan seluruh kota')!='whole_city':raise ValueError('Hanya whole_city yang didukung; batas kecamatan belum dapat ditegakkan')
            policies.append({'kind':'shipping_policy','payload':{'scope':'whole_city','policy':policy,'cities':[text(city,'Kota')],'amount':0,'serviceCoverage':text(coverage,'Layanan'),'costOwner':text(cost,'Penanggung biaya'),'startsAt':timestamp(start,'Mulai'),'endsAt':timestamp(end,'Berakhir'),'approvalRef':text(reference,'Persetujuan'),'verification':{'status':'VERIFIED','evidenceRef':text(evidence,'Bukti'),'verifiedBy':text(approver,'Pemberi persetujuan'),'verifiedAt':timestamp(verified,'Waktu verifikasi')}}})
        return {'sourceSha256':digest,'requests':policies}
    meta={row[0]:row[1] for row in rows(workbook,'Pengesahan Intake')}
    if text(meta.get('status'),'Pengesahan')!='VERIFIED':raise ValueError('Pengesahan harus VERIFIED')
    counted=timestamp(meta.get('countedAt'),'Waktu hitung');purpose=text(meta.get('purpose'),'Tujuan')
    if purpose not in ('initial_inventory','operational_count'):raise ValueError('Tujuan intake tidak sah')
    policies=[];physical={};seen=set()
    for row in rows(workbook,'Intake SKU'):
        sku,status,mode,component,quantity,_proposal,confirmed,version,reference=row[:9]
        if sku not in SKU or sku in seen:raise ValueError('SKU tidak dikenal/duplikat: '+str(sku))
        seen.add(sku)
        if text(status,'SKU '+sku)!='VERIFIED':raise ValueError('SKU harus VERIFIED')
        family,size,*_=SKU[sku]
        if mode not in ('independent','derived') or component not in SKU:raise ValueError('Mode dan SKU komponen belum sah')
        cf,cs,*_=SKU[component];quantity=number(quantity,'Jumlah komponen',1)
        expected=SKU[sku][4]
        if mode=='independent' and (component!=sku or quantity!=1):raise ValueError('SKU independen harus mengonsumsi dirinya sendiri')
        if mode=='derived' and (component!=expected or quantity!=SKU[sku][5]):raise ValueError('Konversi bundle tidak sesuai katalog')
        if mode=='independent':physical[sku]=number(confirmed,'Jumlah fisik '+sku)
        elif confirmed not in (None,0):raise ValueError('Bundle derived tidak memiliki stok independen')
        policies.append({'productId':family,'sizeId':size,'version':number(version,'Versi',1),'approvalRef':text(reference,'Kemasan '+sku),'components':[{'productId':cf,'sizeId':cs,'quantity':quantity}]})
    if len(seen)!=9:raise ValueError('Sembilan kebijakan SKU harus dinyatakan secara eksplisit')
    lots=[];totals={};unique=set()
    for row in rows(workbook,'Intake Lot QC'):
        sku,code,quantity,expected,uom,expiry,count,status,qc,qc_by,verified=row[:11]
        if sku not in physical:raise ValueError('Lot hanya untuk SKU fisik independen')
        if text(verified,'Lot status verifikasi')!='VERIFIED':raise ValueError('Lot harus VERIFIED')
        code=text(code,'Kode lot')
        if status=='released':text(qc_by,'Penanggung jawab QC')
        if (sku,code) in unique:raise ValueError('Lot duplikat')
        unique.add((sku,code))
        if timestamp(count,'Waktu hitung lot')!=counted:raise ValueError('Waktu hitung lot harus sama dengan pengesahan hitung')
        if status not in ('received','quarantined','released','recalled','depleted'):raise ValueError('Status lot tidak sah')
        expiry=text(expiry,'Kedaluwarsa');datetime.date.fromisoformat(expiry)
        quantity=number(quantity,'Jumlah lot');totals[sku]=totals.get(sku,0)+quantity
        family,size,*_=SKU[sku]
        lot={'productId':family,'sizeId':size,'lotCode':code,'unitOfMeasure':text(uom,'UOM'),'expiryDate':expiry,'countedUncommittedQuantity':quantity,'expectedAvailable':number(expected,'Saldo DB'),'status':status}
        if status=='released':lot['qcApprovalRef']=text(qc,'Bukti QC')
        lots.append(lot)
    if not lots:raise ValueError('Lot, expiry, QC dan waktu hitung asli belum tersedia')
    if any(totals.get(sku,0)!=quantity for sku,quantity in physical.items()):raise ValueError('Jumlah lot harus sama dengan alokasi fisik terkonfirmasi')
    verification={'status':'VERIFIED','evidenceRef':text(meta.get('evidenceRef'),'Bukti hitung'),'verifiedBy':text(meta.get('verifiedBy'),'Penanggung jawab'),'verifiedAt':timestamp(meta.get('verifiedAt'),'Waktu verifikasi')}
    inventory={'operationId':'workbook_'+digest[:32],'actor':text(meta.get('actor'),'Akun staf'),'reason':text(meta.get('reason'),'Alasan'),'countedAt':counted,'policies':policies,'lots':lots}
    # Stable physical evidence operation identity; adding formatting cannot duplicate a count.
    inventory['operationId']='count_'+hashlib.sha256(json.dumps({'countedAt':counted,'lots':lots,'policies':policies},sort_keys=True).encode()).hexdigest()[:32]
    return {'documentType':'YUBIE_VERIFIED_INVENTORY_INTAKE_V1','purpose':purpose,'sourceSha256':digest,'verification':verification,'inventory':inventory}

def main():
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('action',choices=['prepare','convert']);parser.add_argument('source',type=pathlib.Path);parser.add_argument('--output',required=True,type=pathlib.Path);parser.add_argument('--scope',choices=['inventory','shipping'],default='inventory');args=parser.parse_args()
    if args.action=='prepare':prepare(args.source,args.output);print('WORKBOOK_PENDING_CREATED_NO_IMPORT');return
    result=convert(args.source,args.scope)
    bridge=pathlib.Path(__file__).with_name('validate-intake.mjs')
    checked=subprocess.run(['node',str(bridge)],input=json.dumps(result),text=True,capture_output=True,timeout=30)
    if checked.returncode:raise ValueError(checked.stderr.strip() or 'Validasi domain/schema gagal')
    with args.output.open('x') as stream:stream.write(json.dumps(result,indent=2)+'\n')
    print('VERIFIED_INTAKE_VALIDATED_NO_DATABASE_MUTATION')

if __name__=='__main__':
    try:main()
    except (ValueError,KeyError,TypeError,zipfile.BadZipFile) as error:print('BLOCKED_INPUT: '+str(error),file=sys.stderr);sys.exit(2)
