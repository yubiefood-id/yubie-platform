import importlib.util, pathlib, tempfile, unittest, datetime, json, subprocess
from openpyxl import Workbook, load_workbook

path = pathlib.Path(__file__).parents[1] / 'workbook_intake.py'
spec = importlib.util.spec_from_file_location('intake', path)
M = importlib.util.module_from_spec(spec)

class IntakeTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        spec.loader.exec_module(M)

    def test_preparation_preserves_original_and_keeps_proposals_pending(self):
        with tempfile.TemporaryDirectory() as tmp:
            source, dest = pathlib.Path(tmp)/'original.xlsx', pathlib.Path(tmp)/'intake.xlsx'
            w=Workbook(); w.active.title='Petunjuk'; w.active['A1']='TEST_ONLY original'; w.save(source)
            original=source.read_bytes(); M.prepare(source,dest)
            self.assertEqual(original,source.read_bytes())
            w=load_workbook(dest); self.assertEqual(w['Petunjuk']['A1'].value,'TEST_ONLY original')
            self.assertEqual(w['Intake SKU']['F2'].value,15)
            self.assertIsNone(w['Intake SKU']['G2'].value)
            with self.assertRaisesRegex(ValueError,'PENDING|VERIFIED'): M.convert(dest,'inventory')

    def test_proposal_json_is_never_an_import(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=pathlib.Path(tmp)/'NOT_FOR_IMPORT.json';p.write_text('{"flour":20}')
            with self.assertRaisesRegex(ValueError,'xlsx'):M.convert(p,'inventory')

    def test_paid_shipping_and_formula_approvals_are_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            source,dest=pathlib.Path(tmp)/'source.xlsx',pathlib.Path(tmp)/'intake.xlsx'
            Workbook().save(source);M.prepare(source,dest);w=load_workbook(dest)
            sheet=w['Intake Pengiriman'];sheet.append(['VERIFIED','TEST_ONLY_CITY','paid',10000,'TEST_ONLY_COURIER','TEST_ONLY_OWNER','2026-01-01T00:00:00Z','2100-01-01T00:00:00Z','TEST_ONLY_DOC','TEST_ONLY_FINANCE','TEST_ONLY_DOC','2026-10-28T00:00:00Z']);w.save(dest)
            with self.assertRaisesRegex(ValueError,'free_promotional'):M.convert(dest,'shipping')
            sheet['C2']='free_promotional';sheet['D2']=0;sheet['I2']='=1+1';w.save(dest)
            with self.assertRaisesRegex(ValueError,'formula'):M.convert(dest,'shipping')

    def test_verified_synthetic_fixture_validates_through_actual_domain_without_database(self):
        with tempfile.TemporaryDirectory() as tmp:
            source,dest=pathlib.Path(tmp)/'source.xlsx',pathlib.Path(tmp)/'intake.xlsx'
            Workbook().save(source);M.prepare(source,dest);w=load_workbook(dest)
            now=datetime.datetime.now(datetime.timezone.utc).isoformat()
            meta={'status':'VERIFIED','evidenceRef':'TEST_ONLY_COUNT','verifiedBy':'TEST_ONLY_COUNTER','verifiedAt':now,'countedAt':now,'purpose':'initial_inventory','actor':'TEST_ONLY_COUNTER','reason':'TEST_ONLY isolated synthetic fixture'}
            for row in w['Pengesahan Intake'].iter_rows(min_row=2):row[1].value=meta[row[0].value]
            for row in w['Intake SKU'].iter_rows(min_row=2):
                row[1].value='VERIFIED';row[8].value='TEST_ONLY_PACKAGING'
                if row[2].value=='independent':
                    row[6].value=row[5].value
                    w['Intake Lot QC'].append([row[0].value,'TEST_ONLY_'+row[0].value,row[5].value,0,row[0].value,'2100-01-01',now,'released','TEST_ONLY_QC','TEST_ONLY_REVIEWER','VERIFIED'])
            w.save(dest);intake=M.convert(dest,'inventory')
            self.assertEqual(sum(l['countedUncommittedQuantity'] for l in intake['inventory']['lots']),80)
            self.assertEqual(len(intake['inventory']['lots']),6)
            bridge=path.parent/'validate-intake.mjs'
            result=subprocess.run(['node',str(bridge)],input=json.dumps(intake),text=True,capture_output=True)
            self.assertEqual(result.returncode,0,result.stderr)
            w['Intake SKU']['G5']=20;w.save(dest)
            with self.assertRaisesRegex(ValueError,'independen'):M.convert(dest,'inventory')

if __name__=='__main__':unittest.main()
